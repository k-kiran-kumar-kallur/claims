"""Output tools: report templating (HTML / PDF / DOCX / JSON), CMS connector, notifications (spec §8)."""
from __future__ import annotations

import io
import json
import os
import smtplib
from email.message import EmailMessage
from pathlib import Path

from fpdf import FPDF
from jinja2 import Environment, FileSystemLoader, select_autoescape
from sqlalchemy.orm import Session

from app.core.logging import get_logger
from app.integrations.cms.base import get_cms_connector
from app.services.events import notify
from app.tools.registry import tool

log = get_logger(__name__)
_TEMPLATES = Path(__file__).resolve().parent.parent / "templates" / "reports"
_env = Environment(loader=FileSystemLoader(str(_TEMPLATES)), autoescape=select_autoescape(["html"]))

_REPL = {"—": "-", "–": "-", "→": "->", "≤": "<=", "≥": ">=", "‘": "'", "’": "'",
         "“": '"', "”": '"', "•": "-", " ": " ", "…": "...", "✓": "v", "✔": "v"}


def _latin1(text: str) -> str:
    for k, v in _REPL.items():
        text = text.replace(k, v)
    text = "".join(ch for ch in text if ch in "\n\t" or ord(ch) >= 32)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def render_html(doc: dict) -> str:
    return _env.get_template("report.html.j2").render(doc=doc)


# Report PDF palette (RGB)
_NAVY, _INK, _MUTED, _RULE = (15, 36, 68), (21, 27, 40), (93, 102, 121), (220, 226, 234)
_ACCENT, _HEAD_BG, _ZEBRA = (31, 95, 191), (238, 242, 247), (248, 250, 252)


class _ReportPDF(FPDF):
    """Letter-size report with a running header (pages 2+) and a confidentiality/page-number footer on every page."""

    def __init__(self, title: str, claim_no: str, generated: str, landscape: bool = False) -> None:
        super().__init__(orientation="L" if landscape else "P", format="Letter", unit="pt")
        self.report_title, self.claim_no, self.generated = title, claim_no, generated
        self.set_margins(48, 54, 48)
        self.set_auto_page_break(True, margin=56)
        self.alias_nb_pages()

    def header(self) -> None:
        if self.page_no() == 1:
            return
        self.set_y(24)
        self.set_font("Helvetica", "B", 7.5)
        self.set_text_color(*_NAVY)
        self.cell(0, 10, _latin1(f"{self.report_title}  |  Claim {self.claim_no}"))
        self.set_font("Helvetica", "", 7.5)
        self.set_text_color(*_MUTED)
        self.set_x(self.l_margin)
        self.cell(0, 10, _latin1(f"Generated {self.generated}"), align="R")
        self.set_draw_color(*_RULE)
        self.set_line_width(0.6)
        self.line(self.l_margin, 38, self.w - self.r_margin, 38)
        self.set_y(self.t_margin)

    def footer(self) -> None:
        self.set_y(-38)
        self.set_draw_color(*_RULE)
        self.set_line_width(0.6)
        self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
        self.ln(5)
        self.set_font("Helvetica", "", 7)
        self.set_text_color(*_MUTED)
        self.cell(0, 9, "CONFIDENTIAL - contains protected health information. Authorized claim-handling use only.")
        self.set_x(self.l_margin)
        self.set_font("Helvetica", "B", 7)
        self.cell(0, 9, f"Page {self.page_no()} of {{nb}}", align="R")


def render_pdf(doc: dict) -> tuple[bytes, int]:
    from datetime import UTC, datetime

    from fpdf.enums import XPos, YPos

    header = doc.get("header") or []
    claim_no = next((str(v) for k, v in header if k.lower().startswith("claim")), "-")
    generated = datetime.now(UTC).strftime("%m/%d/%Y %H:%M UTC")
    pdf = _ReportPDF(doc["title"], claim_no, generated, landscape=bool(doc.get("landscape")))
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin
    left = pdf.l_margin
    nl = {"new_x": XPos.LMARGIN, "new_y": YPos.NEXT, "align": "L"}

    def text_height(w: float, lh: float, text: str) -> float:
        return lh * max(1, len(pdf.multi_cell(w, lh, text, dry_run=True, output="LINES")))

    def para(text: str, size: float = 9, style: str = "", lh: float = 12.5, color=_INK, w: float | None = None, x: float | None = None) -> None:
        pdf.set_font("Helvetica", style, size)
        pdf.set_text_color(*color)
        pdf.set_x(left if x is None else x)
        pdf.multi_cell(w or width, lh, _latin1(text), **nl)
        pdf.set_text_color(*_INK)

    def ensure(space: float) -> None:
        if pdf.get_y() + space > pdf.h - pdf.b_margin:
            pdf.add_page()

    def kv_grid(pairs: list, fill, accent=None) -> None:
        """Two-column grid of label-over-value cells inside a shaded panel."""
        pad, col_gap = 10, 16
        col_w = (width - 2 * pad - col_gap) / 2
        rows = []
        for i in range(0, len(pairs), 2):
            cells = [(str(k), _latin1(str(v if v not in (None, "") else "-"))) for k, v in pairs[i:i + 2]]
            pdf.set_font("Helvetica", "", 9)
            rows.append((cells, 17 + max(text_height(col_w, 11.5, v) for _, v in cells)))
        total = sum(h for _, h in rows) + 2 * pad - 6
        ensure(min(total, 200))
        y0 = pdf.get_y()
        pdf.set_fill_color(*fill)
        pdf.rect(left, y0, width, total, style="F", round_corners=True, corner_radius=4)
        if accent:
            pdf.set_fill_color(*accent)
            pdf.rect(left, y0, 3, total, style="F")
        y = y0 + pad
        for cells, h in rows:
            for j, (k, v) in enumerate(cells):
                x = left + pad + j * (col_w + col_gap)
                pdf.set_xy(x, y)
                pdf.set_font("Helvetica", "B", 6.5)
                pdf.set_text_color(*_MUTED)
                pdf.cell(col_w, 8, _latin1(k.upper()))
                pdf.set_xy(x, y + 9)
                pdf.set_font("Helvetica", "", 9)
                pdf.set_text_color(*_INK)
                pdf.multi_cell(col_w, 11.5, v, new_x=XPos.RIGHT, new_y=YPos.TOP, align="L")
            y += h
        pdf.set_y(y0 + total + 10)

    # ---- title block ----
    pdf.set_fill_color(*_NAVY)
    pdf.rect(0, 0, pdf.w, 6, style="F")
    pdf.set_y(30)
    pdf.set_font("Helvetica", "B", 7.5)
    pdf.set_text_color(*_ACCENT)
    pdf.cell(0, 10, "CLAIMS AI  |  WORKERS' COMPENSATION MEDICAL RECORD REVIEW")
    pdf.set_x(left)
    pdf.set_font("Helvetica", "", 7.5)
    pdf.set_text_color(*_MUTED)
    pdf.cell(0, 10, f"Generated {generated}", align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(8)
    para(doc["title"], 20, "B", 24, _NAVY)
    if doc.get("subtitle"):
        para(doc["subtitle"], 10, "", 14, _MUTED)
    pdf.ln(8)
    if doc.get("banner"):
        pdf.set_font("Helvetica", "B", 8.5)
        h = text_height(width - 24, 12, _latin1(doc["banner"])) + 12
        y0 = pdf.get_y()
        pdf.set_fill_color(255, 248, 225)
        pdf.rect(left, y0, width, h, style="F", round_corners=True, corner_radius=4)
        pdf.set_fill_color(214, 158, 46)
        pdf.rect(left, y0, 3, h, style="F")
        pdf.set_y(y0 + 6)
        para(doc["banner"], 8.5, "B", 12, (122, 83, 0), w=width - 24, x=left + 14)
        pdf.set_y(y0 + h + 10)
    if header:
        kv_grid(header, (245, 247, 251))

    # ---- sections ----
    for n, sec in enumerate(doc.get("sections", []), start=1):
        ensure(140 if sec.get("table", {}).get("rows") or sec.get("summary") else 90)
        pdf.ln(6)
        y = pdf.get_y()
        pdf.set_font("Helvetica", "B", 11.5)
        pdf.set_text_color(*_NAVY)
        pdf.cell(0, 15, _latin1(f"{n}.  {sec['heading']}"), **nl)
        pdf.set_draw_color(*_ACCENT)
        pdf.set_line_width(1.2)
        pdf.line(left, y + 18, left + 28, y + 18)
        pdf.set_draw_color(*_RULE)
        pdf.set_line_width(0.6)
        pdf.line(left + 28, y + 18, left + width, y + 18)
        pdf.set_y(y + 24)
        pdf.set_text_color(*_INK)
        if sec.get("summary"):
            kv_grid(sec["summary"], (234, 241, 252), _ACCENT)
        for label, text in sec.get("fields", []):
            pdf.set_font("Helvetica", "", 9.5)
            ensure(22 + text_height(width, 13, _latin1(str(text))))
            para(label.upper(), 6.5, "B", 9, _MUTED)
            para(str(text), 9.5, "", 13)
            pdf.ln(5)
        if sec.get("note"):
            para(sec["note"], 8.5, "I", 11.5, _MUTED)
            pdf.ln(3)
        for p in sec.get("paragraphs", []):
            para(p, 9.5, "", 13.5)
            pdf.ln(4)
        bullets = sec.get("bullets")
        for b in bullets or []:
            pdf.set_font("Helvetica", "", 9)
            ensure(text_height(width - 14, 12.5, _latin1(b)))
            y = pdf.get_y()
            pdf.set_fill_color(*_ACCENT)
            pdf.ellipse(left + 2, y + 4.5, 3.2, 3.2, style="F")
            para(b, 9, "", 12.5, w=width - 14, x=left + 14)
            pdf.ln(2)
        table = sec.get("table")
        if table is not None and not table.get("rows") and not bullets and not sec.get("paragraphs"):
            para(table.get("empty", "None recorded."), 9, "I", 12, _MUTED)
        if bullets is not None and not bullets:
            para("None recorded.", 9, "I", 12, _MUTED)
        if table and table.get("rows"):
            cols = table["columns"]
            widths = table.get("widths") or [width / len(cols)] * len(cols)
            scale = width / sum(widths)
            widths = [w * scale for w in widths]
            lh, pad = 10.5, 4

            def header_row() -> None:
                pdf.set_font("Helvetica", "B", 7)
                y0 = pdf.get_y()
                hh = max(text_height(w - 2 * pad, 8.5, _latin1(c.upper())) for c, w in zip(cols, widths, strict=False)) + 7
                pdf.set_fill_color(*_HEAD_BG)
                pdf.rect(left, y0, width, hh, style="F")
                x = left
                pdf.set_text_color(*_MUTED)
                for c, w in zip(cols, widths, strict=False):
                    pdf.set_xy(x + pad, y0 + 4)
                    pdf.multi_cell(w - 2 * pad, 8.5, _latin1(c.upper()), new_x=XPos.RIGHT, new_y=YPos.TOP, align="L")
                    x += w
                pdf.set_text_color(*_INK)
                pdf.set_y(y0 + hh)

            ensure(60)
            header_row()
            for i, row in enumerate(table["rows"]):
                cells = [_latin1(str(v if v is not None else "")) for v in row]
                pdf.set_font("Helvetica", "", 8)
                h = max(text_height(w - 2 * pad, lh, t) for t, w in zip(cells, widths, strict=False)) + 2 * pad
                if pdf.get_y() + h > pdf.h - pdf.b_margin:
                    pdf.add_page()
                    header_row()
                    pdf.set_font("Helvetica", "", 8)
                y, x = pdf.get_y(), left
                if i % 2:
                    pdf.set_fill_color(*_ZEBRA)
                    pdf.rect(left, y, width, h, style="F")
                for j, (t, w) in enumerate(zip(cells, widths, strict=False)):
                    pdf.set_font("Helvetica", "B" if j == 0 and table.get("bold_first") else "", 8)
                    pdf.set_xy(x + pad, y + pad)
                    pdf.multi_cell(w - 2 * pad, lh, t, new_x=XPos.RIGHT, new_y=YPos.TOP, align="L")
                    x += w
                pdf.set_draw_color(*_RULE)
                pdf.set_line_width(0.5)
                pdf.line(left, y + h, left + width, y + h)
                pdf.set_xy(left, y + h)
            if table.get("caption"):
                pdf.ln(3)
                para(table["caption"], 7.5, "I", 10, _MUTED)
        pdf.ln(4)

    if doc.get("footer"):
        ensure(50)
        pdf.ln(8)
        pdf.set_draw_color(*_RULE)
        pdf.line(left, pdf.get_y(), left + width, pdf.get_y())
        pdf.ln(6)
        para(doc["footer"], 7.5, "I", 10.5, _MUTED)
    return bytes(pdf.output()), pdf.page_no()


def render_docx(doc: dict) -> bytes:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.size = Pt(9.5)
    d.add_heading(doc["title"], level=1)
    if doc.get("subtitle"):
        d.add_paragraph(doc["subtitle"]).italic = True
    if doc.get("banner"):
        d.add_paragraph(doc["banner"]).runs[0].bold = True
    if doc.get("header"):
        t = d.add_table(rows=0, cols=2)
        t.style = "Light Grid Accent 1"
        for k, v in doc["header"]:
            row = t.add_row().cells
            row[0].text, row[1].text = str(k), str(v or "-")
    for sec in doc.get("sections", []):
        d.add_heading(sec["heading"], level=2)
        if sec.get("summary"):
            t = d.add_table(rows=0, cols=2)
            t.style = "Light Grid Accent 1"
            for k, v in sec["summary"]:
                row = t.add_row().cells
                row[0].text, row[1].text = str(k), str(v or "-")
        for label, text in sec.get("fields", []):
            par = d.add_paragraph()
            par.add_run(f"{label}: ").bold = True
            par.add_run(str(text))
        if sec.get("note"):
            d.add_paragraph(sec["note"]).runs[0].italic = True
        for p in sec.get("paragraphs", []):
            d.add_paragraph(p)
        for b in sec.get("bullets", []):
            d.add_paragraph(b, style="List Bullet")
        table = sec.get("table")
        if table and table.get("rows"):
            t = d.add_table(rows=1, cols=len(table["columns"]))
            t.style = "Light Grid Accent 1"
            for i, c in enumerate(table["columns"]):
                t.rows[0].cells[i].text = c
            for row in table["rows"]:
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = str(v if v is not None else "")
            if table.get("caption"):
                d.add_paragraph(table["caption"]).runs[0].italic = True
        elif table is not None and not sec.get("bullets") and not sec.get("paragraphs"):
            d.add_paragraph(table.get("empty", "None recorded.")).runs[0].italic = True
    if doc.get("footer"):
        d.add_paragraph(doc["footer"]).italic = True
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@tool("report_template_tool", description="Render standardised report templates to HTML, PDF, DOCX or JSON", timeout_s=120)
def report_template_tool(doc: dict, fmt: str) -> dict:
    if fmt == "pdf":
        data, pages = render_pdf(doc)
        return {"bytes": data, "pages": pages, "mime": "application/pdf"}
    if fmt == "docx":
        return {"bytes": render_docx(doc), "pages": None,
                "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
    if fmt == "html":
        return {"bytes": render_html(doc).encode(), "pages": None, "mime": "text/html"}
    return {"bytes": json.dumps(doc, indent=2, default=str).encode(), "pages": None, "mime": "application/json"}


@tool("cms_connector_tool", description="Policy/coverage lookup and posting decision records to the CMS (mock in v1)")
def cms_connector_tool(db: Session, operation: str, claim_id: str, payload: dict | None = None,
                       claim_number: str | None = None) -> dict:
    connector = get_cms_connector()
    if operation == "lookup_policy":
        return connector.lookup_policy(claim_number or "")
    if operation == "post_record":
        return connector.post(db, claim_id, payload or {})
    raise ValueError(f"unknown CMS operation {operation}")


@tool("notification_tool", description="In-app notifications and optional email for assignments and SLA breaches")
def notification_tool(db: Session, user_ids: list[str], kind: str, title: str, body: str | None = None,
                      claim_id: str | None = None, link: str | None = None, emails: list[str] | None = None) -> int:
    for uid in user_ids:
        notify(db, uid, kind, title, body, claim_id, link)
    host = os.environ.get("SMTP_HOST")
    if host and emails:
        try:  # outbox-style best effort; failures never block the workflow
            msg = EmailMessage()
            msg["Subject"] = title
            msg["From"] = os.environ.get("SMTP_FROM", "claims-ai@localhost")
            msg["To"] = ", ".join(emails)
            msg.set_content(f"{body or title}\n\nOpen the Claims AI work queue for details. (No PHI is included in email.)")
            with smtplib.SMTP(host, int(os.environ.get("SMTP_PORT", "25")), timeout=10) as s:
                s.send_message(msg)
        except Exception as exc:
            log.warning("notification_email_failed", error=str(exc))
    return len(user_ids)
